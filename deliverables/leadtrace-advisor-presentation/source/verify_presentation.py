from __future__ import annotations

import sys
import hashlib
import json
import re
import zipfile
from pathlib import Path

import pymupdf
from lxml import html
from PIL import Image, ImageDraw
from pptx import Presentation


SOURCE_DIR = Path(__file__).resolve().parent
ROOT = SOURCE_DIR.parent
RENDERED_DIR = ROOT / "rendered"
HTML_PATH = SOURCE_DIR / "presentation.html"
PPTX_PATH = ROOT / "LeadTrace_导师汇报版.pptx"
PDF_PATH = ROOT / "LeadTrace_导师汇报版.pdf"
PREVIEW_PATH = ROOT / "preview.png"
MANIFEST_PATH = RENDERED_DIR / "manifest.json"

EXPECTED_TITLES = {
    "LeadTrace：把药物化学论文变成可计算的研发知识",
    "AI 提速但没有权力覆盖专家",
    "20篇确定性入库，3个工作区验证端到端闭环",
    "从阅读论文到积累组织研发记忆",
}


def fail(message: str) -> None:
    print(f"FAIL: {message}")
    raise SystemExit(1)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_render_manifest(images: list[Path]) -> None:
    if not MANIFEST_PATH.exists():
        fail("render manifest is missing; rerun render.mjs")
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    if manifest.get("format") != 1:
        fail("render manifest format is unsupported")

    for relative_name, expected_hash in manifest.get("source", {}).items():
        path = SOURCE_DIR / relative_name
        if not path.exists() or sha256(path) != expected_hash:
            fail(f"rendered slides are stale relative to {relative_name}; rerun render.mjs")

    rendered_hashes = manifest.get("rendered", {})
    if set(rendered_hashes) != {path.name for path in images}:
        fail("render manifest does not describe exactly the 12 committed PNGs")
    for image_path in images:
        if sha256(image_path) != rendered_hashes[image_path.name]:
            fail(f"{image_path.name} differs from the render manifest")


def verify_powerpoint_images(images: list[Path]) -> None:
    def media_number(name: str) -> int:
        match = re.search(r"image(\d+)\.png$", name)
        return int(match.group(1)) if match else 10_000

    with zipfile.ZipFile(PPTX_PATH) as archive:
        media_names = sorted(
            (name for name in archive.namelist() if re.search(r"ppt/media/image\d+\.png$", name)),
            key=media_number,
        )
        if len(media_names) != 12:
            fail(f"expected 12 embedded PPTX PNGs, found {len(media_names)}")
        for media_name, image_path in zip(media_names, images, strict=True):
            embedded_hash = hashlib.sha256(archive.read(media_name)).hexdigest()
            if embedded_hash != sha256(image_path):
                fail(f"PPTX media {media_name} is stale relative to {image_path.name}")


def verify_pdf_images(images: list[Path]) -> int:
    pdf = pymupdf.open(PDF_PATH)
    try:
        if pdf.page_count != 12:
            fail(f"expected 12 PDF pages, found {pdf.page_count}")
        for page, image_path in zip(pdf, images, strict=True):
            page_images = page.get_images(full=True)
            if len(page_images) != 1:
                fail(f"PDF page {page.number + 1} does not contain exactly one slide image")
            pixmap = pymupdf.Pixmap(pdf, page_images[0][0])
            with Image.open(image_path) as source_image:
                rgb_image = source_image.convert("RGB")
                if (pixmap.width, pixmap.height) != rgb_image.size or pixmap.samples != rgb_image.tobytes():
                    fail(f"PDF page {page.number + 1} is stale relative to {image_path.name}")
        return pdf.page_count
    finally:
        pdf.close()


def create_contact_sheet(images: list[Path]) -> None:
    thumb_width, thumb_height = 400, 225
    columns, rows = 3, 4
    gutter = 18
    canvas_width = columns * thumb_width + (columns + 1) * gutter
    canvas_height = rows * thumb_height + (rows + 1) * gutter
    sheet = Image.new("RGB", (canvas_width, canvas_height), "#071827")
    draw = ImageDraw.Draw(sheet)

    for index, image_path in enumerate(images):
        with Image.open(image_path) as slide:
            slide = slide.convert("RGB").resize((thumb_width, thumb_height), Image.Resampling.LANCZOS)
            column = index % columns
            row = index // columns
            x = gutter + column * (thumb_width + gutter)
            y = gutter + row * (thumb_height + gutter)
            sheet.paste(slide, (x, y))
            draw.rectangle((x, y, x + thumb_width - 1, y + thumb_height - 1), outline="#355166", width=1)

    sheet.save(PREVIEW_PATH, optimize=True)


def main() -> None:
    if not HTML_PATH.exists():
        fail("presentation.html is missing")

    document = html.parse(str(HTML_PATH))
    slides = document.xpath("//section[contains(concat(' ', normalize-space(@class), ' '), ' slide ')]")
    if len(slides) != 12:
        fail(f"expected 12 HTML slides, found {len(slides)}")

    titles = {slide.get("data-title", "").strip() for slide in slides}
    missing_titles = EXPECTED_TITLES - titles
    if missing_titles:
        fail(f"required titles missing: {sorted(missing_titles)}")

    images = sorted(RENDERED_DIR.glob("[0-9][0-9].png"))
    if len(images) != 12:
        fail(f"expected 12 rendered slides, found {len(images)}")

    for image_path in images:
        with Image.open(image_path) as image:
            if image.size != (1600, 900):
                fail(f"{image_path.name} has dimensions {image.size}, expected 1600x900")

    verify_render_manifest(images)

    if not PPTX_PATH.exists() or PPTX_PATH.stat().st_size < 500_000:
        fail("PowerPoint is missing or unexpectedly small")
    presentation = Presentation(PPTX_PATH)
    if len(presentation.slides) != 12:
        fail(f"expected 12 PowerPoint slides, found {len(presentation.slides)}")
    verify_powerpoint_images(images)

    if not PDF_PATH.exists() or PDF_PATH.stat().st_size < 500_000:
        fail("PDF is missing or unexpectedly small")
    pdf_page_count = verify_pdf_images(images)

    create_contact_sheet(images)
    print("PASS: 12 HTML slides")
    print("PASS: 12 rendered PNGs at 1600x900")
    print("PASS: source hashes match the render manifest")
    print("PASS: 12-slide PowerPoint embeds the current PNGs")
    print(f"PASS: {pdf_page_count}-page PDF embeds the current PNGs")
    print(f"PASS: contact sheet written to {PREVIEW_PATH}")
    print(f"PPTX size: {PPTX_PATH.stat().st_size:,} bytes")
    print(f"PDF size: {PDF_PATH.stat().st_size:,} bytes")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"FAIL: unexpected verification error: {exc}")
        sys.exit(1)
