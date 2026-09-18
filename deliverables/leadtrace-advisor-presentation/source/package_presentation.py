from __future__ import annotations

from pathlib import Path

from lxml import html
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


SOURCE_DIR = Path(__file__).resolve().parent
ROOT = SOURCE_DIR.parent
RENDERED_DIR = ROOT / "rendered"
PPTX_PATH = ROOT / "LeadTrace_导师汇报版.pptx"
PDF_PATH = ROOT / "LeadTrace_导师汇报版.pdf"

SLIDE_WIDTH_IN = 13.333333
SLIDE_HEIGHT_IN = 7.5
PDF_WIDTH_PT = 960
PDF_HEIGHT_PT = 540


def slide_titles() -> list[str]:
    document = html.parse(str(SOURCE_DIR / "presentation.html"))
    return [value.strip() for value in document.xpath("//section[contains(@class, 'slide')]/@data-title")]


def rendered_images() -> list[Path]:
    return sorted(RENDERED_DIR.glob("[0-9][0-9].png"))


def build_powerpoint(images: list[Path], titles: list[str]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = presentation.slide_layouts[6]

    presentation.core_properties.title = "LeadTrace：把药物化学论文变成可计算的研发知识"
    presentation.core_properties.subject = "面向导师的 LeadTrace 项目介绍"
    presentation.core_properties.author = "LeadTrace Project"
    presentation.core_properties.keywords = "LeadTrace, medicinal chemistry, lead optimization, human-in-the-loop"
    presentation.core_properties.comments = "Generated from the repository evidence-backed advisor presentation source."

    for image_path, title in zip(images, titles, strict=True):
        slide = presentation.slides.add_slide(blank_layout)
        picture = slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        picture.name = title
        picture._element.nvPicPr.cNvPr.set("descr", title)

    presentation.save(PPTX_PATH)


def build_pdf(images: list[Path]) -> None:
    pdf = canvas.Canvas(str(PDF_PATH), pagesize=(PDF_WIDTH_PT, PDF_HEIGHT_PT), pageCompression=1)
    pdf.setTitle("LeadTrace：把药物化学论文变成可计算的研发知识")
    pdf.setAuthor("LeadTrace Project")
    pdf.setSubject("面向导师的 LeadTrace 项目介绍")

    for image_path in images:
        pdf.drawImage(
            ImageReader(str(image_path)),
            0,
            0,
            width=PDF_WIDTH_PT,
            height=PDF_HEIGHT_PT,
            preserveAspectRatio=False,
            mask="auto",
        )
        pdf.showPage()
    pdf.save()


def main() -> None:
    images = rendered_images()
    titles = slide_titles()
    if len(images) != 12 or len(titles) != 12:
        raise SystemExit(f"Expected 12 images and titles; found {len(images)} images and {len(titles)} titles")
    build_powerpoint(images, titles)
    build_pdf(images)
    print(f"Created {PPTX_PATH}")
    print(f"Created {PDF_PATH}")


if __name__ == "__main__":
    main()
